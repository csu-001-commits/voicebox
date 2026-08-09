"""Shared styles, palette and source-of-truth data for the Fire Fighting System schedule workbook."""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle

# ----------------------------------------------------------------------------
# BRAND PALETTE (per brief section 15)
# ----------------------------------------------------------------------------
PRIMARY = "4F008C"   # Primary Purple
DEEP    = "2D0054"   # Deep Purple
LIGHT   = "F3E8FF"   # Light Purple
CORAL   = "FF375E"
GREEN   = "059669"
AMBER   = "D97706"
BLUE    = "1D4ED8"
DARK    = "1D252D"   # Dark text
GREY    = "6B7280"
BG      = "F7F4FB"   # Background
WHITE   = "FFFFFF"
RULE    = "E2D9EE"   # hairline separator (derived from light purple)
BAND    = "FBF9FD"   # zebra band

FONT = "Arial"

# ----------------------------------------------------------------------------
# FONTS
# ----------------------------------------------------------------------------
F_TITLE     = Font(name=FONT, size=18, bold=True, color=WHITE)
F_SUBTITLE  = Font(name=FONT, size=10.5, color="E6D9F5")
F_META      = Font(name=FONT, size=9.5, bold=True, color=WHITE)
F_SECTION   = Font(name=FONT, size=11.5, bold=True, color=PRIMARY)
F_HEAD      = Font(name=FONT, size=9, bold=True, color=WHITE)
F_PHASE     = Font(name=FONT, size=9.5, bold=True, color=DEEP)
F_BODY      = Font(name=FONT, size=9, color=DARK)
F_BODY_B    = Font(name=FONT, size=9, bold=True, color=DARK)
F_MUTED     = Font(name=FONT, size=8.5, italic=True, color=GREY)
F_NOTE      = Font(name=FONT, size=9, italic=True, color=DEEP)
F_INPUT     = Font(name=FONT, size=9, color="0000FF")          # editable input cells
F_AUTO      = Font(name=FONT, size=9, bold=True, color=DEEP)   # formula-driven
F_KPI_NUM   = Font(name=FONT, size=22, bold=True, color=PRIMARY)
F_KPI_LBL   = Font(name=FONT, size=8, bold=True, color=GREY)
F_LABEL     = Font(name=FONT, size=9, bold=True, color=DEEP)

# ----------------------------------------------------------------------------
# FILLS
# ----------------------------------------------------------------------------
FILL_HEADER   = PatternFill("solid", fgColor=PRIMARY)
FILL_BANNER   = PatternFill("solid", fgColor=DEEP)
FILL_LIGHT    = PatternFill("solid", fgColor=LIGHT)
FILL_BG       = PatternFill("solid", fgColor=BG)
FILL_WHITE    = PatternFill("solid", fgColor=WHITE)
FILL_BAND     = PatternFill("solid", fgColor=BAND)
FILL_INPUT    = PatternFill("solid", fgColor="FFFDE8")   # cells for the team to fill in
FILL_CORAL    = PatternFill("solid", fgColor="FFE4E9")
FILL_GREEN    = PatternFill("solid", fgColor="D9F2E6")
FILL_AMBER    = PatternFill("solid", fgColor="FDF0DC")
FILL_BLUE     = PatternFill("solid", fgColor="DDE6FB")
FILL_GREY     = PatternFill("solid", fgColor="EFEFF2")

# Gantt bar fills, one per phase
GANTT_FILLS = {
    "Preparation":   PatternFill("solid", fgColor="C9A9E8"),
    "Site Readiness": PatternFill("solid", fgColor="9C6FD1"),
    "Engineering":   PatternFill("solid", fgColor="1D4ED8"),
    "Procurement":   PatternFill("solid", fgColor="6C93F0"),
    "Installation":  PatternFill("solid", fgColor="4F008C"),
    "Testing":       PatternFill("solid", fgColor="D97706"),
    "Handover":      PatternFill("solid", fgColor="059669"),
}

# ----------------------------------------------------------------------------
# BORDERS
# ----------------------------------------------------------------------------
_hair = Side(style="thin", color=RULE)
_med  = Side(style="medium", color=PRIMARY)
_thin_d = Side(style="thin", color="C9B8DC")

B_CELL   = Border(left=_hair, right=_hair, top=_hair, bottom=_hair)
B_HEAD   = Border(left=Side(style="thin", color=PRIMARY), right=Side(style="thin", color=PRIMARY),
                  top=Side(style="thin", color=PRIMARY), bottom=Side(style="medium", color=DEEP))
B_BOX    = Border(left=_thin_d, right=_thin_d, top=_thin_d, bottom=_thin_d)
B_TOP    = Border(top=_med)

# ----------------------------------------------------------------------------
# ALIGNMENTS
# ----------------------------------------------------------------------------
A_WRAP   = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
A_LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
A_CTR    = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_CTR_NW = Alignment(horizontal="center", vertical="center")
A_HEAD   = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ----------------------------------------------------------------------------
# CONTROLLED VOCABULARY (brief section 21)
# ----------------------------------------------------------------------------
STATUS_LIST = ("Not Started,Ready to Start,In Progress,Blocked,At Risk,"
               "Pending Approval,Pending Material,Completed")
PRIORITY_LIST = "Critical,High,Medium,Low"
YESNO_LIST = "Yes,No,TBC"
PROBSEV_LIST = "High,Medium,Low,TBC"
RISKSTATUS_LIST = "Open,Monitoring,Mitigating,Closed"

# RAG fill per status, used by conditional formatting
STATUS_FILLS = [
    ("Completed",        FILL_GREEN,  GREEN),
    ("Blocked",          FILL_CORAL,  CORAL),
    ("At Risk",          FILL_AMBER,  AMBER),
    ("In Progress",      FILL_BLUE,   BLUE),
    ("Pending Approval", FILL_AMBER,  AMBER),
    ("Pending Material", FILL_AMBER,  AMBER),
    ("Ready to Start",   FILL_LIGHT,  PRIMARY),
    ("Not Started",      FILL_GREY,   GREY),
]

# ----------------------------------------------------------------------------
# CONFIRMED SOURCE FACTS — RC Vehicle Survey Report, 30-06-2026
# Nothing outside this block may be presented as confirmed.
# ----------------------------------------------------------------------------
SRC = {
    "project":     "Fire Fighting System for New DC",
    "asset":       "Mobile / Trailer-based DC",
    "survey_date": "30-Jun-2026",
    "pm":          "Kharsan Mujeb Al Ajmi",
    "source_doc":  "RC Vehicle Survey Report – Fire Fighting System for New DC",
    "systems":     ["Fire Alarm System", "Fire Suppression System"],
    "findings": [
        "Installation of Racks and Cabinets has not yet started in the available trailers.",
        "The power source in Trailer-2 is currently unavailable because partition works are still ongoing.",
    ],
}

TBC = "TBC"
DASH = "—"

BASELINE_NOTE = ("PROPOSED PLANNING BASELINE — Proposed planning dates are preliminary and subject to "
                 "confirmation by the Project Owner, PM, Contractor and relevant stakeholders. "
                 "Only items marked CONFIRMED (SURVEY) are supported by the source report.")

LOGO_PLACEHOLDER = "[OFFICIAL LOGO — TO BE INSERTED]"


# ----------------------------------------------------------------------------
# HELPERS
# ----------------------------------------------------------------------------
def set_widths(ws, widths):
    """widths: dict of column letter -> width"""
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def put(ws, coord, value, font=F_BODY, fill=None, align=A_LEFT, border=B_CELL, fmt=None):
    c = ws[coord]
    c.value = value
    c.font = font
    if fill is not None:
        c.fill = fill
    c.alignment = align
    if border is not None:
        c.border = border
    if fmt:
        c.number_format = fmt
    return c


def banner(ws, row, last_col, title, subtitle, meta_left, meta_right):
    """Corporate document header band spanning row..row+2."""
    ws.merge_cells(f"A{row}:{last_col}{row}")
    ws.merge_cells(f"A{row+1}:{last_col}{row+1}")
    ws.merge_cells(f"A{row+2}:{last_col}{row+2}")

    ws.row_dimensions[row].height = 30
    ws.row_dimensions[row + 1].height = 18
    ws.row_dimensions[row + 2].height = 16

    t = ws[f"A{row}"]
    t.value = title
    t.font = F_TITLE
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    s = ws[f"A{row+1}"]
    s.value = subtitle
    s.font = F_SUBTITLE
    s.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    m = ws[f"A{row+2}"]
    m.value = f"{meta_left}          |          {meta_right}          |          {LOGO_PLACEHOLDER}"
    m.font = F_META
    m.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    from openpyxl.utils import column_index_from_string
    last = column_index_from_string(last_col)
    for r in (row, row + 1, row + 2):
        for ci in range(1, last + 1):
            ws.cell(row=r, column=ci).fill = FILL_BANNER if r > row else FILL_HEADER


def note_bar(ws, row, last_col, text, fill=FILL_LIGHT, font=F_NOTE, height=26):
    ws.merge_cells(f"A{row}:{last_col}{row}")
    ws.row_dimensions[row].height = height
    c = ws[f"A{row}"]
    c.value = text
    c.font = font
    c.fill = fill
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    from openpyxl.utils import column_index_from_string
    for ci in range(1, column_index_from_string(last_col) + 1):
        ws.cell(row=row, column=ci).fill = fill
        ws.cell(row=row, column=ci).border = B_BOX


def section_title(ws, row, last_col, text):
    ws.merge_cells(f"A{row}:{last_col}{row}")
    ws.row_dimensions[row].height = 22
    c = ws[f"A{row}"]
    c.value = text
    c.font = F_SECTION
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.border = Border(bottom=Side(style="medium", color=PRIMARY))
    from openpyxl.utils import column_index_from_string
    for ci in range(1, column_index_from_string(last_col) + 1):
        ws.cell(row=row, column=ci).border = Border(bottom=Side(style="medium", color=PRIMARY))


def table_header(ws, row, headers, start_col=1, height=32):
    ws.row_dimensions[row].height = height
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i)
        c.value = h
        c.font = F_HEAD
        c.fill = FILL_HEADER
        c.alignment = A_HEAD
        c.border = B_HEAD


def page_setup(ws, last_col, title_rows=None, landscape=True, fit_width=1):
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if landscape else ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = fit_width
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.6
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddFooter.left.text = "Confidential — Project Planning Document"
    ws.oddFooter.left.size = 8
    ws.oddFooter.left.color = "6B7280"
    ws.oddFooter.center.text = ("Prepared based on the RC Vehicle Survey Report dated 30-Jun-2026  |  "
                                "Version 0.1 — Preliminary Planning Baseline")
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = "6B7280"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.color = "6B7280"
    ws.sheet_view.showGridLines = False
