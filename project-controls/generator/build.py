#!/usr/bin/env python3
"""Build the Fire Fighting System — New DC installation schedule workbook."""

import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter, column_index_from_string
from datetime import datetime

from common import *          # noqa
from data import *            # noqa

OUT = sys.argv[1] if len(sys.argv) > 1 else "output.xlsx"

DATEFMT = "DD-MMM-YYYY"

SH_DASH = "1. Exec Dashboard"
SH_SCHED = "2. Master Schedule"
SH_GANTT = "3. Gantt Chart"
SH_RESP = "4. Responsibility Matrix"
SH_ACT = "5. Action Tracker"
SH_RISK = "6. Risk Register"
SH_DOC = "7. Document Control"

TITLE = "FIRE FIGHTING SYSTEM — NEW DC"
SUBTITLE = "Installation Schedule & Project Responsibility Tracker"
META_L = f"Survey Date: {SRC['survey_date']}"
META_R = f"PM: {SRC['pm']}"


# ---------------------------------------------------------------------------
# generic helpers
# ---------------------------------------------------------------------------
def mrange(ws, ref, value=None, font=F_BODY, fill=None, align=A_LEFT, border=B_CELL, fmt=None):
    if ":" in ref:
        ws.merge_cells(ref)
        cells = [c for row in ws[ref] for c in row]
    else:
        cells = [ws[ref]]
    for cell in cells:
        cell.font = font
        if fill is not None:
            cell.fill = fill
        if border is not None:
            cell.border = border
    anchor = ws[ref.split(":")[0]]
    anchor.value = value
    anchor.alignment = align
    if fmt:
        anchor.number_format = fmt
    return anchor


def status_cf(ws, ref):
    """Apply RAG conditional formatting to a range of Status cells."""
    for text, fill, colour in STATUS_FILLS:
        ws.conditional_formatting.add(
            ref,
            CellIsRule(operator="equal", formula=[f'"{text}"'],
                       fill=fill, font=Font(name=FONT, size=9, bold=True, color=colour)),
        )


def add_dv(ws, ranges, values, prompt_title, prompt):
    dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True,
                        showDropDown=False, errorStyle="warning")
    dv.promptTitle = prompt_title
    dv.prompt = prompt
    dv.showInputMessage = True
    dv.showErrorMessage = True
    dv.error = "Please choose a value from the drop-down list."
    dv.errorTitle = "Value not in list"
    ws.add_data_validation(dv)
    for r in ranges:
        dv.add(r)
    return dv


wb = Workbook()

# ===========================================================================
# SHEET 2 — MASTER SCHEDULE  (built first: the dashboard aggregates from it)
# ===========================================================================
ws = wb.active
ws.title = SH_SCHED
LAST = "U"

SCHED_COLS = [
    ("A", 6,  "ID"),
    ("B", 8,  "WBS"),
    ("C", 26, "Phase"),
    ("D", 48, "Activity"),
    ("E", 19, "Trailer /\nLocation"),
    ("F", 11, "Prede-\ncessor"),
    ("G", 15, "Planned Start\n(enter date)"),
    ("H", 15, "Planned Finish\n(enter date)"),
    ("I", 10, "Duration\n(days)"),
    ("J", 17, "Status"),
    ("K", 11, "%\nComplete"),
    ("L", 22, "Responsible\nPerson"),
    ("M", 26, "Role"),
    ("N", 15, "Project\nOwner"),
    ("O", 18, "Contractor"),
    ("P", 34, "Dependency"),
    ("Q", 34, "Risk / Blocker"),
    ("R", 12, "Approval\nRequired"),
    ("S", 52, "Remarks"),
    ("T", 12, "Days\nRemaining\n(auto)"),
    ("U", 30, "Schedule Flag (auto)"),
]
set_widths(ws, {c: w for c, w, _ in SCHED_COLS})

banner(ws, 1, LAST, TITLE, SUBTITLE, META_L, META_R)
ws.row_dimensions[4].height = 6
note_bar(ws, 5, LAST, BASELINE_NOTE, fill=FILL_LIGHT, height=30)
note_bar(ws, 6, LAST,
         "HOW TO USE  ·  Shaded input cells (Planned Start / Planned Finish / % Complete) and blue text are for the "
         "project team to complete.  ·  Duration, Days Remaining and Schedule Flag calculate automatically.  ·  "
         "Status is a drop-down list.  ·  Remarks tagged “CONFIRMED (SURVEY)” are supported by the source report; "
         "everything else is a proposed planning value or TBC.",
         fill=FILL_BG, font=F_MUTED, height=28)
ws.row_dimensions[7].height = 6

HEADER_ROW = 8
table_header(ws, HEADER_ROW, [h for _, _, h in SCHED_COLS], height=40)

row = HEADER_ROW + 1
data_rows = []          # row numbers holding activities
phase_rows = []

current_phase = None
band = False
for a in ACTIVITIES:
    (aid, pkey, activity, location, pred, status, pct, resp, role,
     contractor, dependency, risk, approval, remarks, wks, wke, flag) = a

    phase_label, phase_group = PHASES[pkey]

    if pkey != current_phase:
        # phase banner row
        env = [x for x in ACTIVITIES if x[1] == pkey and x[14] > 0]
        wmin = min(x[14] for x in env) if env else 0
        wmax = max(x[15] for x in env) if env else 0
        span = f"Proposed baseline: Week {wmin} – Week {wmax}" if wmin else "Completed prior to baseline"
        mrange(ws, f"A{row}:{LAST}{row}", f"   {phase_label}          ·          {span}",
               font=F_PHASE, fill=FILL_LIGHT, align=A_LEFT,
               border=Border(top=Side(style="thin", color=PRIMARY),
                             bottom=Side(style="thin", color=PRIMARY)))
        ws.row_dimensions[row].height = 20
        phase_rows.append(row)
        row += 1
        current_phase = pkey
        band = False

    fill = FILL_BAND if band else FILL_WHITE
    band = not band

    wbs_root = pkey[1] + ".0"
    vals = [
        len(data_rows) + 1, aid, phase_label, activity, location, pred,
        None, None, None, status, pct, resp, role, TBC, contractor,
        dependency, risk, approval, remarks, None, None,
    ]
    for i, v in enumerate(vals):
        col = get_column_letter(i + 1)
        c = ws[f"{col}{row}"]
        c.value = v
        c.fill = fill
        c.border = B_CELL
        c.font = F_BODY
        c.alignment = A_WRAP

    ws[f"B{row}"].value = aid
    ws[f"C{row}"].value = phase_label

    # alignment / typography refinements
    for col in ("A", "B", "F", "I", "K", "R", "T"):
        ws[f"{col}{row}"].alignment = A_CTR
    for col in ("G", "H", "J", "N", "U"):
        ws[f"{col}{row}"].alignment = A_CTR
    ws[f"D{row}"].font = F_BODY_B
    ws[f"S{row}"].font = F_MUTED if not remarks.startswith(CONFIRMED) else Font(
        name=FONT, size=8.5, bold=True, color=DEEP)

    # ---- date cells: user input (only 1.02 carries a source-confirmed date) ----
    for col in ("G", "H"):
        c = ws[f"{col}{row}"]
        c.number_format = DATEFMT
        c.font = F_INPUT
        c.fill = FILL_INPUT
    if aid == "1.02":
        d = datetime(2026, 6, 30)
        ws[f"G{row}"].value = d
        ws[f"H{row}"].value = d
        ws[f"G{row}"].font = F_BODY_B
        ws[f"H{row}"].font = F_BODY_B
        ws[f"G{row}"].fill = fill
        ws[f"H{row}"].fill = fill

    # ---- % complete: user input ----
    k = ws[f"K{row}"]
    k.number_format = "0%"
    k.font = F_BODY_B if pct else F_INPUT
    if not pct:
        k.fill = FILL_INPUT

    # ---- formulas ----
    ws[f"I{row}"] = f'=IF(OR($G{row}="",$H{row}=""),"TBD",$H{row}-$G{row}+1)'
    ws[f"I{row}"].font = F_AUTO
    ws[f"T{row}"] = (f'=IF($H{row}="","TBD",IF($J{row}="Completed","Closed",$H{row}-TODAY()))')
    ws[f"T{row}"].font = F_AUTO
    ws[f"U{row}"] = (
        f'=IF($J{row}="Completed","COMPLETE",'
        f'IF(AND($H{row}<>"",$H{row}<TODAY()),"OVERDUE",'
        f'IF($J{row}="Blocked","BLOCKED",'
        f'IF(ISNUMBER(SEARCH("Power availability",$P{row})),"PENDING POWER AVAILABILITY",'
        f'IF(IFERROR(IF($F{row}="{DASH}",FALSE,'
        f'INDEX($J:$J,MATCH($F{row},$B:$B,0))<>"Completed"),FALSE),"PENDING PREDECESSOR",'
        f'IF($J{row}="At Risk","AT RISK",'
        f'IF(OR($J{row}="Pending Approval",$J{row}="Pending Material"),UPPER($J{row}),'
        f'IF($J{row}="In Progress","IN PROGRESS",'
        f'IF($J{row}="Ready to Start","READY TO START","ON PLAN"))))))))'
    )
    ws[f"U{row}"].font = F_AUTO

    # ---- flag highlighting ----
    if flag == "T2 POWER":
        for col in ("D", "P", "Q"):
            ws[f"{col}{row}"].fill = FILL_CORAL
    elif flag == "CRITICAL":
        ws[f"D{row}"].fill = FILL_LIGHT

    data_rows.append(row)
    row += 1

FIRST_DATA, LAST_DATA = data_rows[0], data_rows[-1]
END = row - 1

# conditional formatting -----------------------------------------------------
status_cf(ws, f"J{FIRST_DATA}:J{LAST_DATA}")

flag_rules = [
    ("OVERDUE", CORAL, "FFE4E9"), ("BLOCKED", CORAL, "FFE4E9"),
    ("PENDING POWER AVAILABILITY", CORAL, "FFE4E9"),
    ("PENDING PREDECESSOR", AMBER, "FDF0DC"), ("AT RISK", AMBER, "FDF0DC"),
    ("PENDING APPROVAL", AMBER, "FDF0DC"), ("PENDING MATERIAL", AMBER, "FDF0DC"),
    ("COMPLETE", GREEN, "D9F2E6"), ("IN PROGRESS", BLUE, "DDE6FB"),
    ("READY TO START", PRIMARY, "F3E8FF"), ("ON PLAN", GREY, "EFEFF2"),
]
for text, colour, bgc in flag_rules:
    ws.conditional_formatting.add(
        f"U{FIRST_DATA}:U{LAST_DATA}",
        CellIsRule(operator="equal", formula=[f'"{text}"'],
                   fill=PatternFill("solid", fgColor=bgc),
                   font=Font(name=FONT, size=8.5, bold=True, color=colour)))

# overdue highlight across the row
ws.conditional_formatting.add(
    f"A{FIRST_DATA}:U{LAST_DATA}",
    FormulaRule(formula=[f'AND($H{FIRST_DATA}<>"",$H{FIRST_DATA}<TODAY(),$J{FIRST_DATA}<>"Completed")'],
                fill=PatternFill("solid", fgColor="FFF0F3"), stopIfTrue=False))

# data validation ------------------------------------------------------------
groups, start = [], data_rows[0]
prev = data_rows[0]
for r in data_rows[1:]:
    if r != prev + 1:
        groups.append((start, prev)); start = r
    prev = r
groups.append((start, prev))

add_dv(ws, [f"J{a}:J{b}" for a, b in groups], STATUS_LIST,
       "Status", "Select the current activity status.")
add_dv(ws, [f"R{a}:R{b}" for a, b in groups], YESNO_LIST,
       "Approval Required", "Does this activity require a formal approval?")

ws.freeze_panes = f"E{HEADER_ROW+1}"
ws.auto_filter.ref = f"A{HEADER_ROW}:U{LAST_DATA}"
# 21 columns cannot print legibly on a single sheet width: allow two page-widths
# and repeat the identifying columns on the second.
# No repeat-title ROW: the activity table fits one page height, and a repeated header
# would otherwise print on top of the Critical Dependencies table below it.
page_setup(ws, LAST, fit_width=2)
ws.print_title_cols = "A:D"

# --- critical dependencies block (brief section 8) --------------------------
# start it on a fresh printed page so it carries its own header
from openpyxl.worksheet.pagebreak import Break
ws.row_breaks.append(Break(id=END + 1))

r = END + 2
section_title(ws, r, LAST, "CRITICAL DEPENDENCIES & CONSTRAINTS  —  confirmed by the RC Vehicle Survey Report")
r += 2
dep_heads = ["Ref", "Confirmed Constraint", "Reason", "Dependency", "Required Action",
             "Responsible Party", "Target Resolution", "Impact on Schedule", "Classification"]
dep_spans = ["A:A", "B:D", "E:F", "G:I", "J:L", "M:O", "P:P", "Q:S", "T:U"]
ws.row_dimensions[r].height = 30
for h, sp in zip(dep_heads, dep_spans):
    c1, c2 = sp.split(":")
    mrange(ws, f"{c1}{r}:{c2}{r}", h, font=F_HEAD, fill=FILL_HEADER, align=A_HEAD, border=B_HEAD)
r += 1
for d in DEPENDENCIES:
    ws.row_dimensions[r].height = 58
    for val, sp in zip(d, dep_spans):
        c1, c2 = sp.split(":")
        f = F_BODY_B if sp in ("A:A", "B:D", "T:U") else F_BODY
        fl = FILL_CORAL if sp in ("A:A", "T:U") else FILL_WHITE
        mrange(ws, f"{c1}{r}:{c2}{r}", val, font=f, fill=fl, align=A_WRAP)
    r += 1

note_bar(ws, r + 1, LAST,
         "Both constraints above are stated in the source survey report. Target resolution dates are TBD until "
         "confirmed by the Project Owner and the responsible contractors.",
         fill=FILL_BG, font=F_MUTED, height=22)

SCHED_STATUS_RANGE = f"'{SH_SCHED}'!$J${FIRST_DATA}:$J${LAST_DATA}"
SCHED_PCT_RANGE = f"'{SH_SCHED}'!$K${FIRST_DATA}:$K${LAST_DATA}"
SCHED_WBS_RANGE = f"'{SH_SCHED}'!$B${FIRST_DATA}:$B${LAST_DATA}"
SCHED_FLAG_RANGE = f"'{SH_SCHED}'!$U${FIRST_DATA}:$U${LAST_DATA}"


# ===========================================================================
# SHEET 3 — GANTT CHART
# ===========================================================================
gw = wb.create_sheet(SH_GANTT)
NWEEKS = 18
WK_FIRST_COL = 9                       # column I
WK_LAST_COL = WK_FIRST_COL + NWEEKS - 1
GLAST = get_column_letter(WK_LAST_COL)

set_widths(gw, {"A": 6, "B": 8, "C": 52, "D": 17, "E": 15, "F": 11, "G": 10, "H": 10})
for i in range(WK_FIRST_COL, WK_LAST_COL + 1):
    gw.column_dimensions[get_column_letter(i)].width = 4.4

banner(gw, 1, GLAST, TITLE, SUBTITLE, META_L, META_R)
gw.row_dimensions[4].height = 6
note_bar(gw, 5, GLAST,
         "PROPOSED BASELINE — DATES TO BE CONFIRMED.  The timeline below is relative (Week 1 = project kickoff). "
         "It is NOT a calendar schedule. Proposed planning dates are preliminary and subject to confirmation by the "
         "Project Owner, PM, Contractor and relevant stakeholders.",
         fill=FILL_CORAL, font=Font(name=FONT, size=9, bold=True, color=DEEP), height=30)

# legend — three items per row across the week columns, six columns per item
LEG = [
    (GANTT_FILLS["Preparation"],    "Preparation"),
    (GANTT_FILLS["Site Readiness"], "Site Readiness"),
    (GANTT_FILLS["Engineering"],    "Engineering"),
    (GANTT_FILLS["Installation"],   "Installation"),
    (GANTT_FILLS["Testing"],        "Testing & Commissioning"),
    (GANTT_FILLS["Handover"],       "Handover"),
    (PatternFill("solid", fgColor=DEEP),  "Critical activity"),
    (PatternFill("solid", fgColor=CORAL), "Trailer-2 power dependency"),
]
LEG_TOP, LEG_ROWS = 6, 3
mrange(gw, f"A{LEG_TOP}:H{LEG_TOP+LEG_ROWS-1}", "LEGEND", font=F_LABEL, fill=FILL_BG, align=A_CTR)
for i, (fill, label) in enumerate(LEG):
    lr = LEG_TOP + i // 3
    slot = i % 3
    sc = WK_FIRST_COL + slot * 6
    gw.row_dimensions[lr].height = 15
    swatch = gw.cell(row=lr, column=sc)
    swatch.fill = fill
    swatch.border = B_CELL
    mrange(gw, f"{get_column_letter(sc+1)}{lr}:{get_column_letter(sc+5)}{lr}", label,
           font=F_MUTED, fill=FILL_BG, align=A_LEFT, border=None)
for lr in range(LEG_TOP, LEG_TOP + LEG_ROWS):
    for ci in range(WK_FIRST_COL, WK_LAST_COL + 1):
        if gw.cell(row=lr, column=ci).fill.fgColor.rgb in (None, "00000000"):
            gw.cell(row=lr, column=ci).fill = FILL_BG
gw.row_dimensions[LEG_TOP + LEG_ROWS].height = 6

GH = LEG_TOP + LEG_ROWS + 1
gw.row_dimensions[GH].height = 34
for i, h in enumerate(["ID", "WBS", "Activity", "Phase", "Critical /\nDependency",
                       "Prede-\ncessor", "Baseline\nWk Start", "Baseline\nWk Finish"]):
    c = gw.cell(row=GH, column=i + 1)
    c.value = h
    c.font = F_HEAD
    c.fill = FILL_HEADER
    c.alignment = A_HEAD
    c.border = B_HEAD
for w in range(1, NWEEKS + 1):
    c = gw.cell(row=GH, column=WK_FIRST_COL + w - 1)
    c.value = w
    c.number_format = '"W"0'
    c.font = Font(name=FONT, size=8, bold=True, color=WHITE)
    c.fill = FILL_HEADER
    c.alignment = A_CTR_NW
    c.border = B_HEAD

grow = GH + 1
gdata_first = grow
current_phase = None
seq = 0
for a in ACTIVITIES:
    (aid, pkey, activity, location, pred, status, pct, resp, role,
     contractor, dependency, risk, approval, remarks, wks, wke, flag) = a
    phase_label, phase_group = PHASES[pkey]

    if pkey != current_phase:
        env = [x for x in ACTIVITIES if x[1] == pkey and x[14] > 0]
        wmin = min(x[14] for x in env) if env else 0
        wmax = max(x[15] for x in env) if env else 0
        gw.row_dimensions[grow].height = 20
        mrange(gw, f"A{grow}:C{grow}", f"   {phase_label}", font=F_PHASE, fill=FILL_LIGHT,
               align=A_LEFT, border=B_CELL)
        for cl, v, f in (("D", phase_group, F_PHASE), ("E", "PHASE ENVELOPE", F_MUTED),
                         ("F", DASH, F_PHASE), ("G", wmin or None, F_PHASE), ("H", wmax or None, F_PHASE)):
            c = gw[f"{cl}{grow}"]
            c.value = v
            c.font = f
            c.fill = FILL_LIGHT
            c.border = B_CELL
            c.alignment = A_CTR
        for i in range(WK_FIRST_COL, WK_LAST_COL + 1):
            gw.cell(row=grow, column=i).border = B_CELL
        current_phase = pkey
        grow += 1

    seq += 1
    gw.row_dimensions[grow].height = 15
    row_vals = [seq, aid, activity, phase_group,
                flag if flag else DASH, pred, wks or None, wke or None]
    for i, v in enumerate(row_vals):
        c = gw.cell(row=grow, column=i + 1)
        c.value = v
        c.border = B_CELL
        c.font = F_BODY_B if i == 2 else F_BODY
        c.alignment = A_LEFT if i == 2 else A_CTR
    if flag == "T2 POWER":
        gw[f"E{grow}"].font = Font(name=FONT, size=8.5, bold=True, color=CORAL)
        gw[f"C{grow}"].fill = FILL_CORAL
    elif flag == "CRITICAL":
        gw[f"E{grow}"].font = Font(name=FONT, size=8.5, bold=True, color=PRIMARY)
    else:
        gw[f"E{grow}"].font = F_MUTED
    for i in range(WK_FIRST_COL, WK_LAST_COL + 1):
        gw.cell(row=grow, column=i).border = B_CELL
    grow += 1

gdata_last = grow - 1

# planned completion marker row
grow += 1
mrange(gw, f"A{grow}:F{grow}", "   PLANNED PROJECT COMPLETION — HANDOVER & CLOSEOUT (proposed)",
       font=Font(name=FONT, size=9.5, bold=True, color=WHITE), fill=FILL_BANNER, align=A_LEFT)
for cl, v in (("G", NWEEKS), ("H", NWEEKS)):
    c = gw[f"{cl}{grow}"]
    c.value = v
    c.font = Font(name=FONT, size=9, bold=True, color=WHITE)
    c.fill = FILL_BANNER
    c.alignment = A_CTR
    c.border = B_CELL
for i in range(WK_FIRST_COL, WK_LAST_COL + 1):
    c = gw.cell(row=grow, column=i)
    c.border = B_CELL
    if i == WK_LAST_COL:
        c.value = "◆"
        c.font = Font(name=FONT, size=11, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=GREEN)
        c.alignment = A_CTR_NW
    else:
        c.fill = FILL_BG
gw.row_dimensions[grow].height = 20
MARKER_ROW = grow

BARS = f"{get_column_letter(WK_FIRST_COL)}{gdata_first}:{GLAST}{gdata_last}"
FC = get_column_letter(WK_FIRST_COL)

# priority order: Trailer-2 dependency, then critical, then phase colour
gw.conditional_formatting.add(BARS, FormulaRule(
    formula=[f'AND($E{gdata_first}="T2 POWER",{FC}${GH}>=$G{gdata_first},{FC}${GH}<=$H{gdata_first})'],
    fill=PatternFill("solid", fgColor=CORAL), stopIfTrue=True))
gw.conditional_formatting.add(BARS, FormulaRule(
    formula=[f'AND($E{gdata_first}="CRITICAL",{FC}${GH}>=$G{gdata_first},{FC}${GH}<=$H{gdata_first})'],
    fill=PatternFill("solid", fgColor=DEEP), stopIfTrue=True))
for key, fill in GANTT_FILLS.items():
    gw.conditional_formatting.add(BARS, FormulaRule(
        formula=[f'AND($D{gdata_first}="{key}",{FC}${GH}>=$G{gdata_first},{FC}${GH}<=$H{gdata_first})'],
        fill=fill, stopIfTrue=True))
# phase envelope bars
ENV = f"{FC}{GH+1}:{GLAST}{gdata_last}"
for key, fill in GANTT_FILLS.items():
    gw.conditional_formatting.add(ENV, FormulaRule(
        formula=[f'AND($E{GH+1}="PHASE ENVELOPE",$D{GH+1}="{key}",{FC}${GH}>=$G{GH+1},{FC}${GH}<=$H{GH+1})'],
        fill=fill, stopIfTrue=True))

gw.freeze_panes = f"{FC}{GH+1}"
page_setup(gw, GLAST, title_rows=f"{GH}:{GH}")

note_bar(gw, MARKER_ROW + 2, GLAST,
         "Dependencies are shown by the Predecessor column; each activity may only start once its predecessor is "
         "complete. Activities marked CRITICAL sit on the critical path. Activities marked T2 POWER cannot proceed "
         "until partition works in Trailer-2 are complete and the power source is confirmed available.",
         fill=FILL_BG, font=F_MUTED, height=32)


# ===========================================================================
# SHEET 4 — RESPONSIBILITY MATRIX
# ===========================================================================
rw = wb.create_sheet(SH_RESP)
RLAST = "G"
set_widths(rw, {"A": 32, "B": 28, "C": 30, "D": 28, "E": 62, "F": 24, "G": 18})
banner(rw, 1, RLAST, TITLE, SUBTITLE, META_L, META_R)
rw.row_dimensions[4].height = 6

note_bar(rw, 5, RLAST,
         "The only person identified in the source survey report is Kharsan Mujeb Al Ajmi (Project Manager). "
         "All other names are marked TBC — To Be Confirmed. Names and contact details are editable fields for the "
         "project team to complete; no names have been assumed.",
         fill=FILL_LIGHT, height=30)

section_title(rw, 7, RLAST, "PROJECT GOVERNANCE  —  editable")
GOV = [("Project Owner", TBC), ("Project Manager", SRC["pm"]), ("Main Contractor", TBC),
       ("Fire System Contractor", TBC), ("Client Representative", TBC), ("QA/QC", TBC), ("HSE", TBC)]
r = 8
for label, val in GOV:
    rw.row_dimensions[r].height = 20
    mrange(rw, f"A{r}", label, font=F_LABEL, fill=FILL_LIGHT, align=A_LEFT)
    confirmed = val != TBC
    mrange(rw, f"B{r}:C{r}", val,
           font=F_BODY_B if confirmed else F_INPUT,
           fill=FILL_WHITE if confirmed else FILL_INPUT, align=A_LEFT)
    mrange(rw, f"D{r}:G{r}",
           "Confirmed by the source survey report." if confirmed
           else "Enter the confirmed name here once appointed.",
           font=F_MUTED, fill=FILL_WHITE, align=A_LEFT)
    r += 1

r += 1
section_title(rw, r, RLAST, "PROJECT RESPONSIBILITY MATRIX")
r += 2
RESP_HEADS = ["Responsibility Area", "Name", "Role", "Organization / Contractor",
              "Responsibility", "Contact", "Status"]
table_header(rw, r, RESP_HEADS, height=30)
RESP_H = r
r += 1
resp_first = r
for i, item in enumerate(RESPONSIBILITY):
    rw.row_dimensions[r].height = 30
    fill = FILL_BAND if i % 2 else FILL_WHITE
    for j, v in enumerate(item):
        c = rw.cell(row=r, column=j + 1)
        c.value = v
        c.border = B_CELL
        c.fill = fill
        c.alignment = A_WRAP if j in (4,) else A_LEFT
        if j == 0:
            c.font = F_BODY_B
        elif j in (1, 5) and v == TBC:
            c.font = F_INPUT
            c.fill = FILL_INPUT
        elif j == 1:
            c.font = F_BODY_B
        else:
            c.font = F_BODY
        if j in (3,) and v == TBC:
            c.font = F_INPUT
            c.fill = FILL_INPUT
        if j == 6:
            c.alignment = A_CTR
    r += 1
resp_last = r - 1
status_cf(rw, f"G{resp_first}:G{resp_last}")
add_dv(rw, [f"G{resp_first}:G{resp_last}"], STATUS_LIST, "Status", "Select the current status.")
rw.freeze_panes = f"A{RESP_H+1}"
page_setup(rw, RLAST, title_rows=f"{RESP_H}:{RESP_H}")
note_bar(rw, resp_last + 2, RLAST,
         "TBC — To Be Confirmed.  Shaded cells are editable fields for the project team.",
         fill=FILL_BG, font=F_MUTED, height=20)


# ===========================================================================
# SHEET 5 — ACTION TRACKER
# ===========================================================================
aw = wb.create_sheet(SH_ACT)
ALAST = "I"
set_widths(aw, {"A": 10, "B": 50, "C": 52, "D": 28, "E": 15, "F": 12, "G": 17, "H": 13, "I": 46})
banner(aw, 1, ALAST, TITLE, SUBTITLE, META_L, META_R)
aw.row_dimensions[4].height = 6
note_bar(aw, 5, ALAST,
         "MANAGEMENT ACTION TRACKER — actions ordered by their ability to delay installation. "
         "Critical path sequence: complete partition works in Trailer-2 → confirm power availability → "
         "complete racks and cabinets → confirm site readiness → release installation activities. "
         "Due dates are TBD until the baseline is confirmed.",
         fill=FILL_LIGHT, height=32)
AH = 7
table_header(aw, AH, ["Action ID", "Required Action", "Reason", "Responsible", "Due Date",
                      "Priority", "Status", "Escalation\nRequired", "Remarks"], height=32)
r = AH + 1
act_first = r
for i, item in enumerate(ACTIONS):
    aw.row_dimensions[r].height = 34
    fill = FILL_BAND if i % 2 else FILL_WHITE
    for j, v in enumerate(item):
        c = aw.cell(row=r, column=j + 1)
        c.value = v
        c.border = B_CELL
        c.fill = fill
        c.font = F_BODY
        c.alignment = A_WRAP
        if j == 0:
            c.font = F_BODY_B
            c.alignment = A_CTR
        if j == 1:
            c.font = F_BODY_B
        if j in (4, 5, 6, 7):
            c.alignment = A_CTR
        if j == 4:
            c.font = F_INPUT
            c.fill = FILL_INPUT
        if j == 8:
            c.font = F_MUTED
    if item[5] == "Critical":
        aw[f"A{r}"].fill = FILL_CORAL
        aw[f"B{r}"].fill = FILL_CORAL
    r += 1
act_last = r - 1
status_cf(aw, f"G{act_first}:G{act_last}")
for text, colour, bgc in (("Critical", CORAL, "FFE4E9"), ("High", AMBER, "FDF0DC"),
                          ("Medium", BLUE, "DDE6FB"), ("Low", GREY, "EFEFF2")):
    aw.conditional_formatting.add(f"F{act_first}:F{act_last}", CellIsRule(
        operator="equal", formula=[f'"{text}"'], fill=PatternFill("solid", fgColor=bgc),
        font=Font(name=FONT, size=9, bold=True, color=colour)))
add_dv(aw, [f"G{act_first}:G{act_last}"], STATUS_LIST, "Status", "Select the current action status.")
add_dv(aw, [f"F{act_first}:F{act_last}"], PRIORITY_LIST, "Priority", "Select the action priority.")
add_dv(aw, [f"H{act_first}:H{act_last}"], YESNO_LIST, "Escalation", "Does this action require escalation?")
aw.freeze_panes = f"A{AH+1}"
page_setup(aw, ALAST, title_rows=f"{AH}:{AH}")
note_bar(aw, act_last + 2, ALAST,
         "Actions A01–A04 form the critical path to releasing installation. Shaded Due Date cells are editable.",
         fill=FILL_BG, font=F_MUTED, height=20)


# ===========================================================================
# SHEET 6 — RISK REGISTER
# ===========================================================================
kw = wb.create_sheet(SH_RISK)
KLAST = "L"
set_widths(kw, {"A": 9, "B": 42, "C": 46, "D": 50, "E": 14, "F": 12, "G": 56,
                "H": 30, "I": 15, "J": 14, "K": 20, "L": 30})
banner(kw, 1, KLAST, TITLE, SUBTITLE, META_L, META_R)
kw.row_dimensions[4].height = 6
note_bar(kw, 5, KLAST,
         "Risks R01 and R02 are the constraints explicitly recorded in the source survey report. All remaining "
         "entries are standard project-planning risks and are labelled “Planning Risk — To Be Validated”. "
         "Probability, Severity and Status are drop-down lists; the Risk Rating column calculates automatically.",
         fill=FILL_LIGHT, height=32)
KH = 7
table_header(kw, KH, ["Risk ID", "Risk", "Cause", "Impact", "Probability", "Severity",
                      "Mitigation", "Owner", "Target Date", "Status",
                      "Risk Rating\n(auto)", "Classification"], height=32)
r = KH + 1
risk_first = r
for i, item in enumerate(RISKS):
    kw.row_dimensions[r].height = 42
    fill = FILL_BAND if i % 2 else FILL_WHITE
    rid, risk, cause, impact, prob, sev, mit, owner, tgt, st, cls = item
    vals = [rid, risk, cause, impact, prob, sev, mit, owner, tgt, st, None, cls]
    for j, v in enumerate(vals):
        c = kw.cell(row=r, column=j + 1)
        c.value = v
        c.border = B_CELL
        c.fill = fill
        c.font = F_BODY
        c.alignment = A_WRAP
        if j == 0:
            c.font = F_BODY_B
            c.alignment = A_CTR
        if j == 1:
            c.font = F_BODY_B
        if j in (4, 5, 8, 9, 10):
            c.alignment = A_CTR
        if j == 8:
            c.font = F_INPUT
            c.fill = FILL_INPUT
        if j == 11:
            c.font = Font(name=FONT, size=8, bold=True,
                          color=CORAL if v == CONF_RISK else GREY)
            c.alignment = A_CTR
    kw[f"K{r}"] = (
        f'=IF(OR($E{r}="",$F{r}="",$E{r}="TBC",$F{r}="TBC"),"TBC",'
        f'IF(AND($E{r}="High",$F{r}="High"),"CRITICAL",'
        f'IF(OR(AND($E{r}="High",$F{r}="Medium"),AND($E{r}="Medium",$F{r}="High")),"HIGH",'
        f'IF(AND($E{r}="Low",$F{r}="Low"),"LOW","MEDIUM"))))'
    )
    kw[f"K{r}"].font = F_AUTO
    if cls == CONF_RISK:
        kw[f"A{r}"].fill = FILL_CORAL
        kw[f"B{r}"].fill = FILL_CORAL
    r += 1
risk_last = r - 1
for text, colour, bgc in (("CRITICAL", CORAL, "FFE4E9"), ("HIGH", CORAL, "FFE4E9"),
                          ("MEDIUM", AMBER, "FDF0DC"), ("LOW", GREEN, "D9F2E6")):
    kw.conditional_formatting.add(f"K{risk_first}:K{risk_last}", CellIsRule(
        operator="equal", formula=[f'"{text}"'], fill=PatternFill("solid", fgColor=bgc),
        font=Font(name=FONT, size=9, bold=True, color=colour)))
for col in ("E", "F"):
    for text, colour, bgc in (("High", CORAL, "FFE4E9"), ("Medium", AMBER, "FDF0DC"),
                              ("Low", GREEN, "D9F2E6")):
        kw.conditional_formatting.add(f"{col}{risk_first}:{col}{risk_last}", CellIsRule(
            operator="equal", formula=[f'"{text}"'], fill=PatternFill("solid", fgColor=bgc),
            font=Font(name=FONT, size=9, bold=True, color=colour)))
add_dv(kw, [f"E{risk_first}:F{risk_last}"], PROBSEV_LIST, "Probability / Severity", "High, Medium or Low.")
add_dv(kw, [f"J{risk_first}:J{risk_last}"], RISKSTATUS_LIST, "Risk Status", "Select the current risk status.")
kw.freeze_panes = f"B{KH+1}"
page_setup(kw, KLAST, title_rows=f"{KH}:{KH}")
note_bar(kw, risk_last + 2, KLAST,
         "Risk Rating is derived from Probability × Severity: High/High = CRITICAL; High/Medium = HIGH; "
         "Low/Low = LOW; all other combinations = MEDIUM.",
         fill=FILL_BG, font=F_MUTED, height=20)


# ===========================================================================
# SHEET 7 — DOCUMENT CONTROL
# ===========================================================================
dw = wb.create_sheet(SH_DOC)
DLAST = "F"
set_widths(dw, {"A": 32, "B": 62, "C": 20, "D": 22, "E": 26, "F": 46})
banner(dw, 1, DLAST, TITLE, SUBTITLE, META_L, META_R)
dw.row_dimensions[4].height = 6

section_title(dw, 6, DLAST, "DOCUMENT CONTROL")
DOCINFO = [
    ("Document Title", "Fire Fighting System Installation Schedule", False),
    ("Project", "New DC", False),
    ("Asset / Location", SRC["asset"], False),
    ("Source Document", SRC["source_doc"], False),
    ("Survey Date", SRC["survey_date"], False),
    ("Project Manager", SRC["pm"], False),
    ("Version", "0.1", False),
    ("Status", "Preliminary — Planning Baseline", False),
    ("Prepared By", TBC, True),
    ("Reviewed By", TBC, True),
    ("Approved By", TBC, True),
    ("Last Updated", TBC, True),
]
r = 7
for label, val, editable in DOCINFO:
    dw.row_dimensions[r].height = 22
    mrange(dw, f"A{r}", label, font=F_LABEL, fill=FILL_LIGHT, align=A_LEFT)
    mrange(dw, f"B{r}:F{r}", val,
           font=F_INPUT if editable else F_BODY_B,
           fill=FILL_INPUT if editable else FILL_WHITE, align=A_LEFT)
    r += 1

r += 1
section_title(dw, r, DLAST, "REVISION HISTORY")
r += 2
table_header(dw, r, ["Version", "Description of Change", "Date", "Prepared By", "Reviewed By", "Approved By"],
             height=26)
r += 1
REV = [("0.1", "Initial preliminary planning baseline prepared from the RC Vehicle Survey Report "
               "dated 30-Jun-2026. All dates proposed and subject to confirmation.",
        TBC, TBC, TBC, TBC)]
for i, item in enumerate(REV):
    dw.row_dimensions[r].height = 34
    for j, v in enumerate(item):
        c = dw.cell(row=r, column=j + 1)
        c.value = v
        c.border = B_CELL
        c.fill = FILL_WHITE
        c.font = F_BODY_B if j == 0 else (F_INPUT if v == TBC else F_BODY)
        if v == TBC:
            c.fill = FILL_INPUT
        c.alignment = A_WRAP if j == 1 else A_CTR
    r += 1
for extra in range(4):
    dw.row_dimensions[r].height = 20
    for j in range(6):
        c = dw.cell(row=r, column=j + 1)
        c.border = B_CELL
        c.fill = FILL_INPUT
        c.font = F_INPUT
        c.alignment = A_CTR
    r += 1

r += 1
section_title(dw, r, DLAST, "SOURCE TRACEABILITY  —  what is confirmed and what is proposed")
r += 2
table_header(dw, r, ["Item", "Value / Statement", "Classification", "Source", "Sheet Reference", "Note"], height=26)
r += 1
TRACE = [
    ("Project", SRC["project"], "CONFIRMED", SRC["source_doc"], "All sheets", "Stated in the survey report."),
    ("Location / Asset", SRC["asset"], "CONFIRMED", SRC["source_doc"], "All sheets", "Stated in the survey report."),
    ("Survey Date", SRC["survey_date"], "CONFIRMED", SRC["source_doc"], "All sheets", "Stated in the survey report."),
    ("Project Manager", SRC["pm"], "CONFIRMED", SRC["source_doc"], "All sheets",
     "The only person named in the source."),
    ("Site condition", SRC["findings"][0], "CONFIRMED", SRC["source_doc"], "2. Master Schedule (2.04) / CD-02",
     "Drives the rack and cabinet readiness dependency."),
    ("Site condition", SRC["findings"][1], "CONFIRMED", SRC["source_doc"], "2. Master Schedule (2.02) / CD-01",
     "Drives the Trailer-2 power dependency."),
    ("Recommended system", "Fire Alarm System", "CONFIRMED", SRC["source_doc"], "2. Master Schedule (Phase 3–5)",
     "Recommended by the survey."),
    ("Recommended system", "Fire Suppression System", "CONFIRMED", SRC["source_doc"], "2. Master Schedule (Phase 3–5)",
     "Recommended by the survey."),
    ("Start / finish dates", "Not stated in the source", "PROPOSED / TBD", "Not available", "All sheets",
     "Relative Week 1–18 baseline only; calendar dates TBD."),
    ("Overall duration", "Not stated in the source", "PROPOSED / TBD", "Not available", "3. Gantt Chart",
     "18-week proposed baseline, subject to confirmation."),
    ("Project Owner / contractors", "Not stated in the source", "TBC", "Not available", "4. Responsibility Matrix",
     "No names assumed."),
    ("Equipment types", "Not stated in the source", "TBC", "Not available", "2. Master Schedule (Phase 3–4)",
     "To be confirmed during engineering/design."),
]
for i, item in enumerate(TRACE):
    dw.row_dimensions[r].height = 30
    fill = FILL_BAND if i % 2 else FILL_WHITE
    for j, v in enumerate(item):
        c = dw.cell(row=r, column=j + 1)
        c.value = v
        c.border = B_CELL
        c.fill = fill
        c.font = F_BODY_B if j == 0 else F_BODY
        c.alignment = A_WRAP
        if j == 2:
            c.alignment = A_CTR
            c.font = Font(name=FONT, size=8.5, bold=True,
                          color=GREEN if v == "CONFIRMED" else AMBER)
        if j == 5:
            c.font = F_MUTED
    r += 1

note_bar(dw, r + 1, DLAST,
         "Confidential — Project Planning Document.  Prepared based on the RC Vehicle Survey Report dated "
         "30-Jun-2026.  Version 0.1 — Preliminary Planning Baseline.  Last updated: TBC.",
         fill=FILL_LIGHT, height=24)
page_setup(dw, DLAST, landscape=False)


# ===========================================================================
# SHEET 1 — EXECUTIVE DASHBOARD  (built last, moved to front)
# ===========================================================================
ew = wb.create_sheet(SH_DASH)
ELAST = "S"
ew.column_dimensions["A"].width = 2.2
for i in range(2, 20):
    ew.column_dimensions[get_column_letter(i)].width = 8.6

banner(ew, 1, ELAST, TITLE, SUBTITLE, META_L, META_R)
ew.row_dimensions[4].height = 8

note_bar(ew, 5, ELAST, BASELINE_NOTE, fill=FILL_LIGHT, height=30)
ew.row_dimensions[6].height = 8

# ---- Executive summary -----------------------------------------------------
section_title(ew, 7, ELAST, "EXECUTIVE SUMMARY")
r = 8
SUMMARY = [
    ("Project", SRC["project"]),
    ("Location / Asset", SRC["asset"]),
    ("Survey Date", SRC["survey_date"]),
    ("Project Manager", SRC["pm"]),
    ("Source Document", SRC["source_doc"]),
    ("Recommended Systems", "Fire Alarm System  ·  Fire Suppression System"),
]
for label, val in SUMMARY:
    ew.row_dimensions[r].height = 19
    mrange(ew, f"B{r}:E{r}", label, font=F_LABEL, fill=FILL_LIGHT, align=A_LEFT)
    mrange(ew, f"F{r}:S{r}", val, font=F_BODY_B, fill=FILL_WHITE, align=A_LEFT)
    r += 1

ew.row_dimensions[r].height = 6
r += 1
ew.row_dimensions[r].height = 44
mrange(ew, f"B{r}:E{r}", "Current Situation", font=F_LABEL, fill=FILL_LIGHT, align=A_LEFT)
mrange(ew, f"F{r}:S{r}",
       "Per the RC Vehicle Survey Report dated 30-Jun-2026: installation of Racks and Cabinets has not yet started "
       "in the available trailers, and the power source in Trailer-2 is currently unavailable because partition "
       "works are still ongoing. The survey recommends providing a Fire Alarm System and a Fire Suppression System. "
       "No installation start, duration or completion dates are stated in the source document.",
       font=F_BODY, fill=FILL_WHITE, align=A_WRAP)
r += 1
ew.row_dimensions[r].height = 26
mrange(ew, f"B{r}:E{r}", "Overall Status", font=F_LABEL, fill=FILL_LIGHT, align=A_LEFT)
mrange(ew, f"F{r}:S{r}", "PLANNING / SITE READINESS PENDING",
       font=Font(name=FONT, size=12, bold=True, color=CORAL), fill=FILL_CORAL, align=A_LEFT)
r += 2

# ---- KPI tiles -------------------------------------------------------------
section_title(ew, r, ELAST, "KEY PERFORMANCE INDICATORS  —  calculated live from the Master Schedule")
r += 2
KPI_ROW = r
TILES = [
    ("B", "D", "TOTAL ACTIVITIES", f"=COUNTA({SCHED_WBS_RANGE})", "0", "Across 6 phases", PRIMARY),
    ("E", "G", "COMPLETED", f'=COUNTIF({SCHED_STATUS_RANGE},"Completed")', "0", "Confirmed by survey", GREEN),
    ("H", "J", "IN PROGRESS", f'=COUNTIF({SCHED_STATUS_RANGE},"In Progress")', "0", "Partition works ongoing", BLUE),
    ("K", "M", "BLOCKED", f'=COUNTIF({SCHED_STATUS_RANGE},"Blocked")', "0", "Trailer-2 power", CORAL),
    ("N", "P", "NOT STARTED", f'=COUNTIF({SCHED_STATUS_RANGE},"Not Started")', "0", "Awaiting site readiness", GREY),
    ("Q", "S", "OVERALL COMPLETE", f'=IFERROR(AVERAGE({SCHED_PCT_RANGE}),0)', "0.0%", "Weighted equally", PRIMARY),
]
ew.row_dimensions[KPI_ROW].height = 16
ew.row_dimensions[KPI_ROW + 1].height = 32
ew.row_dimensions[KPI_ROW + 2].height = 16
for c1, c2, label, formula, fmt, sub, colour in TILES:
    mrange(ew, f"{c1}{KPI_ROW}:{c2}{KPI_ROW}", label, font=F_KPI_LBL, fill=FILL_BG,
           align=A_CTR, border=Border(left=Side(style="thin", color=RULE),
                                      right=Side(style="thin", color=RULE),
                                      top=Side(style="medium", color=colour)))
    mrange(ew, f"{c1}{KPI_ROW+1}:{c2}{KPI_ROW+1}", formula,
           font=Font(name=FONT, size=20, bold=True, color=colour), fill=FILL_BG,
           align=A_CTR_NW, fmt=fmt,
           border=Border(left=Side(style="thin", color=RULE), right=Side(style="thin", color=RULE)))
    mrange(ew, f"{c1}{KPI_ROW+2}:{c2}{KPI_ROW+2}", sub, font=F_KPI_LBL, fill=FILL_BG, align=A_CTR,
           border=Border(left=Side(style="thin", color=RULE), right=Side(style="thin", color=RULE),
                         bottom=Side(style="thin", color=RULE)))
r = KPI_ROW + 4

# ---- Critical constraints --------------------------------------------------
section_title(ew, r, ELAST, "CRITICAL CONSTRAINTS  —  confirmed by the survey")
r += 2
CARDS = [
    ("B", "J", "CONSTRAINT 01  ·  RACKS & CABINETS",
     "Installation of Racks and Cabinets has not yet started in the available trailers.",
     "Impact: gates fire alarm and suppression activities that interface with racks and cabinets "
     "(4.01, 4.04, 4.07) and therefore the whole installation phase.\n"
     "Required action: confirm scope, contractor and start date (Action A03).\n"
     "Responsible: Main Contractor (TBC)      Target resolution: TBD"),
    ("K", "S", "CONSTRAINT 02  ·  TRAILER-2 POWER",
     "The power source in Trailer-2 is currently unavailable because partition works are still ongoing.",
     "Impact: blocks power availability verification (2.02) and power and control connections (4.08); "
     "gates functional testing (5.05) and Trailer-2 handover.\n"
     "Required action: complete partition works, then verify and energise (Actions A01, A02).\n"
     "Responsible: Civil / Partition Contractor (TBC) → Electrical Contractor (TBC)      Target resolution: TBD"),
]
ew.row_dimensions[r].height = 20
ew.row_dimensions[r + 1].height = 32
ew.row_dimensions[r + 2].height = 62
for c1, c2, head, statement, detail in CARDS:
    mrange(ew, f"{c1}{r}:{c2}{r}", head, font=Font(name=FONT, size=9.5, bold=True, color=WHITE),
           fill=PatternFill("solid", fgColor=CORAL), align=A_LEFT)
    mrange(ew, f"{c1}{r+1}:{c2}{r+1}", statement, font=F_BODY_B, fill=FILL_CORAL, align=A_WRAP)
    mrange(ew, f"{c1}{r+2}:{c2}{r+2}", detail, font=F_BODY, fill=FILL_WHITE, align=A_WRAP)
r += 4

# ---- Milestones ------------------------------------------------------------
section_title(ew, r, ELAST, "KEY MILESTONES")
r += 2
MSPANS = [("B", "E"), ("F", "G"), ("H", "I"), ("J", "K"), ("L", "N"), ("O", "S")]
MHEADS = ["Milestone", "Target Date", "Owner", "Status", "Dependency", "Remarks"]
ew.row_dimensions[r].height = 26
for (c1, c2), h in zip(MSPANS, MHEADS):
    mrange(ew, f"{c1}{r}:{c2}{r}", h, font=F_HEAD, fill=FILL_HEADER, align=A_HEAD, border=B_HEAD)
MH = r
r += 1
ms_first = r
for i, m in enumerate(MILESTONES):
    _, name, tgt, owner, status, dep, remarks = m
    ew.row_dimensions[r].height = 26
    fill = FILL_BAND if i % 2 else FILL_WHITE
    vals = [name, tgt, owner, status, dep, remarks]
    for (c1, c2), v in zip(MSPANS, vals):
        editable = v in (TBC, "TBD")
        mrange(ew, f"{c1}{r}:{c2}{r}", v,
               font=(F_BODY_B if c1 == "B" else (F_INPUT if editable else F_BODY)),
               fill=(FILL_INPUT if editable else fill),
               align=(A_LEFT if c1 in ("B", "L", "O") else A_CTR))
    if "CONFIRMED constraint" in remarks:
        mrange(ew, f"B{r}:E{r}", name, font=Font(name=FONT, size=9, bold=True, color=CORAL),
               fill=FILL_CORAL, align=A_LEFT)
    r += 1
ms_last = r - 1
status_cf(ew, f"J{ms_first}:K{ms_last}")
add_dv(ew, [f"J{ms_first}:J{ms_last}"], STATUS_LIST, "Status", "Select the milestone status.")
r += 1

# ---- Management attention --------------------------------------------------
section_title(ew, r, ELAST, "MANAGEMENT ATTENTION  —  critical path actions")
r += 2
ASPANS = [("B", "C"), ("D", "I"), ("J", "L"), ("M", "N"), ("O", "P"), ("Q", "S")]
AHEADS = ["Action ID", "Required Action", "Responsible", "Priority", "Status", "Escalation / Remarks"]
ew.row_dimensions[r].height = 26
for (c1, c2), h in zip(ASPANS, AHEADS):
    mrange(ew, f"{c1}{r}:{c2}{r}", h, font=F_HEAD, fill=FILL_HEADER, align=A_HEAD, border=B_HEAD)
r += 1
top_first = r
for i, a in enumerate([x for x in ACTIONS if x[5] == "Critical"] + [x for x in ACTIONS if x[5] == "High"][:2]):
    aid, action, reason, resp, due, prio, st, esc, rem = a
    ew.row_dimensions[r].height = 26
    fill = FILL_BAND if i % 2 else FILL_WHITE
    vals = [aid, action, resp, prio, st, rem]
    for (c1, c2), v in zip(ASPANS, vals):
        mrange(ew, f"{c1}{r}:{c2}{r}", v,
               font=F_BODY_B if c1 in ("B", "D") else F_BODY, fill=fill,
               align=A_LEFT if c1 in ("D", "J", "Q") else A_CTR)
    r += 1
top_last = r - 1
status_cf(ew, f"O{top_first}:P{top_last}")
for text, colour, bgc in (("Critical", CORAL, "FFE4E9"), ("High", AMBER, "FDF0DC")):
    ew.conditional_formatting.add(f"M{top_first}:N{top_last}", CellIsRule(
        operator="equal", formula=[f'"{text}"'], fill=PatternFill("solid", fgColor=bgc),
        font=Font(name=FONT, size=9, bold=True, color=colour)))
r += 1

# ---- Legend ----------------------------------------------------------------
section_title(ew, r, ELAST, "STATUS KEY & DOCUMENT LEGEND")
r += 2
LEGEND = [
    ("🟢", "Completed / On Track", GREEN, "🟡", "At Risk / Pending", AMBER),
    ("🔴", "Blocked / Critical", CORAL, "⚪", "Not Started", GREY),
    ("🔵", "In Progress", BLUE, "▨", "Shaded cell = editable field for the project team", DEEP),
]
for a1, t1, c1c, a2, t2, c2c in LEGEND:
    ew.row_dimensions[r].height = 18
    mrange(ew, f"B{r}", a1, font=Font(name=FONT, size=11, color=c1c), fill=FILL_WHITE, align=A_CTR)
    mrange(ew, f"C{r}:I{r}", t1, font=F_BODY, fill=FILL_WHITE, align=A_LEFT)
    mrange(ew, f"J{r}", a2, font=Font(name=FONT, size=11, color=c2c), fill=FILL_WHITE, align=A_CTR)
    mrange(ew, f"K{r}:S{r}", t2, font=F_BODY, fill=FILL_WHITE, align=A_LEFT)
    r += 1
r += 1
note_bar(ew, r, ELAST,
         "CONFIRMED INFORMATION FROM SURVEY is limited to: project and location, survey date 30-Jun-2026, "
         "Project Manager Kharsan Mujeb Al Ajmi, the two site constraints above, and the recommendation to provide "
         "a Fire Alarm System and a Fire Suppression System. Everything else in this workbook is a PROPOSED "
         "PLANNING BASELINE or marked TBC / TBD.",
         fill=FILL_BG, font=F_MUTED, height=32)
r += 1
note_bar(ew, r, ELAST,
         "Confidential — Project Planning Document  ·  Prepared based on the RC Vehicle Survey Report dated "
         "30-Jun-2026  ·  Version 0.1 — Preliminary Planning Baseline  ·  Last updated: TBC",
         fill=FILL_LIGHT, font=Font(name=FONT, size=8.5, bold=True, color=DEEP), height=22)

page_setup(ew, ELAST)

# reorder: dashboard first
wb.move_sheet(SH_DASH, offset=-(len(wb.sheetnames) - 1))
wb.active = 0
for name in wb.sheetnames:
    wb[name].sheet_properties.tabColor = PRIMARY

wb.save(OUT)
print(f"saved {OUT}")
print(f"master schedule data rows: {FIRST_DATA}-{LAST_DATA} ({len(data_rows)} activities)")
print(f"gantt rows: {gdata_first}-{gdata_last}")
